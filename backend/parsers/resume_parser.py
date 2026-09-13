import os
import io
import re
from typing import Tuple

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    import docx
except ImportError:
    docx = None


def extract_text_from_pdf_bytes(content_bytes: bytes) -> str:
    """Extract text from PDF file bytes using PyPDF."""
    if not PdfReader:
        raise ImportError("pypdf library is not installed.")
    
    text_chunks = []
    pdf_file = io.BytesIO(content_bytes)
    reader = PdfReader(pdf_file)
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text_chunks.append(page_text)
            
    return "\n".join(text_chunks)


def extract_text_from_docx_bytes(content_bytes: bytes) -> str:
    """Extract text from DOCX file bytes using python-docx."""
    if not docx:
        raise ImportError("python-docx library is not installed.")
    
    docx_file = io.BytesIO(content_bytes)
    doc = docx.Document(docx_file)
    text_chunks = [paragraph.text for paragraph in doc.paragraphs if paragraph.text]
    
    # Also extract table text if present
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                text_chunks.append(row_text)
                
    return "\n".join(text_chunks)


def extract_text_from_txt_bytes(content_bytes: bytes) -> str:
    """Extract text from UTF-8 / ASCII text file bytes."""
    try:
        return content_bytes.decode("utf-8")
    except UnicodeDecodeError:
        return content_bytes.decode("latin-1", errors="ignore")


def parse_resume_file(filename: str, content_bytes: bytes) -> Tuple[str, str]:
    """
    Parses incoming resume file based on extension (.pdf, .docx, .txt).
    Returns (cleaned_text, format_type).
    """
    ext = os.path.splitext(filename)[1].lower()
    
    if ext == ".pdf":
        raw_text = extract_text_from_pdf_bytes(content_bytes)
        fmt = "PDF"
    elif ext in [".docx", ".doc"]:
        raw_text = extract_text_from_docx_bytes(content_bytes)
        fmt = "DOCX"
    elif ext in [".txt", ".md"]:
        raw_text = extract_text_from_txt_bytes(content_bytes)
        fmt = "TXT"
    else:
        # Fallback decoding attempt
        raw_text = extract_text_from_txt_bytes(content_bytes)
        fmt = "UNKNOWN/TXT"
        
    cleaned = clean_raw_extracted_text(raw_text)
    return cleaned, fmt


def clean_raw_extracted_text(text: str) -> str:
    """Basic text cleanup for raw extracted resume documents."""
    if not text:
        return ""
    # Normalize line breaks and tabs
    text = re.sub(r'\r\n|\r', '\n', text)
    text = re.sub(r'\t+', ' ', text)
    # Remove multiple space padding while keeping structure
    lines = [line.strip() for line in text.split('\n')]
    # Remove empty lines excess
    cleaned_lines = [line for line in lines if line]
    return "\n".join(cleaned_lines)


def anonymize_resume_text(text: str, candidate_name: str = None) -> str:
    """
    Redacts Personally Identifiable Information (PII) from resume text
    to support ethical, bias-free candidate evaluation.
    """
    if not text:
        return ""
        
    anonymized = text
    
    # 1. Redact Candidate Name if provided
    if candidate_name and candidate_name.strip():
        name_parts = [p.strip() for p in candidate_name.split() if len(p.strip()) > 1]
        for part in name_parts:
            pattern = re.compile(re.escape(part), re.IGNORECASE)
            anonymized = pattern.sub("[CANDIDATE NAME REDACTED]", anonymized)
            
    # 2. Redact Email addresses
    anonymized = re.sub(
        r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}',
        '[EMAIL REDACTED]',
        anonymized
    )
    
    # 3. Redact Phone numbers (US, international formats)
    anonymized = re.sub(
        r'\(?\b\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}\b',
        '[PHONE REDACTED]',
        anonymized
    )
    anonymized = re.sub(
        r'\+?\d{1,3}[-.\s]?\(?\d{2,4}\)?[-.\s]?\d{3,4}[-.\s]?\d{3,4}',
        '[PHONE REDACTED]',
        anonymized
    )
    
    # 4. Redact Web links & LinkedIn/GitHub profiles
    anonymized = re.sub(
        r'https?://[^\s]+|www\.[^\s]+|linkedin\.com/in/[^\s]+|github\.com/[^\s]+',
        '[LINK REDACTED]',
        anonymized,
        flags=re.IGNORECASE
    )
    
    # 5. Redact header candidate name heuristic (first non-empty line if short)
    lines = anonymized.split('\n')
    if lines and len(lines[0].split()) <= 4 and '[CANDIDATE' not in lines[0]:
        lines[0] = "[CANDIDATE ANONYMIZED]"
        anonymized = "\n".join(lines)
        
    return anonymized

